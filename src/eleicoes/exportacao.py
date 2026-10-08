"""Planilhas para a equipe de rua. Números continuam numéricos, para ordenar no Excel."""

from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl.utils import get_column_letter

from eleicoes.formatacao import fmt_turno
from eleicoes.normalizacao import titulo_pt


def _sim_nao(serie: pd.Series) -> pd.Series:
    return serie.map(lambda valor: "Sim" if bool(valor) else "Não")


def _bairros(df: pd.DataFrame) -> pd.DataFrame:
    dados = {
        "Bairro": df["bairro"],
        "Vencedor": df.apply(_rotulo_vencedor, axis=1),
        "Percentual do vencedor": df["percentual_vencedor"],
        "Segundo colocado": df["segundo"],
        "Margem (p.p.)": df["margem_pp"],
        "Votos válidos": df["votos_validos"],
        "Votos nominais": df["votos_nominais"],
        "Votos de legenda": df["votos_legenda"],
        "Brancos": df["brancos"],
        "Nulos": df["nulos"],
        "Aptos": df["aptos"],
        "Comparecimento": df["comparecimento"],
        "Abstenção": df["abstencao"],
        "Abstenção (%)": df["abstencao_pct"],
        "Prioritário": _sim_nao(df["prioritario"]),
        "Seções": df["secoes"],
        "Locais de votação": df["locais"],
    }
    if "lider_disputa" in df.columns:
        dados["Na frente no 2º turno"] = df["lider_disputa"]
        dados["Adversário no 2º turno"] = df["adversario_disputa"]
        dados["Votos de quem está na frente"] = df["votos_lider_disputa"]
        dados["Votos do adversário"] = df["votos_adversario_disputa"]
        dados["Margem da disputa (p.p.)"] = df["margem_disputa_pp"]
        dados["Votos em aberto"] = df["votos_outros"]
        dados["Pode virar"] = _sim_nao(df["pode_virar"])
        dados["Prioritário no 2º turno"] = _sim_nao(df["prioritario_2t"])
    return pd.DataFrame(dados)


def _rotulo_vencedor(linha: pd.Series) -> str:
    if bool(linha.get("empate")) and linha.get("empate_entre"):
        return f"Empate entre {linha['empate_entre']}"
    return str(linha.get("vencedor", ""))


def _candidatos(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return pd.DataFrame(columns=["Bairro", "Candidato", "Número", "Tipo", "Votos", "Percentual", "Posição"])
    tipo = df["tipo"].map({"nominal": "Nominal", "legenda": "Legenda"}).fillna(df["tipo"])
    return pd.DataFrame(
        {
            "Bairro": df["bairro"],
            "Candidato": df["candidato"],
            "Número": df["nr_votavel"],
            "Tipo": tipo,
            "Votos": df["votos"],
            "Percentual": df["percentual"],
            "Posição": df["posicao"],
        }
    )


def _locais(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return pd.DataFrame()
    return pd.DataFrame(
        {
            "Bairro": df["bairro"],
            "Local de votação": df["local_votacao"],
            "Endereço": df["endereco"],
            "Latitude": df["latitude"],
            "Longitude": df["longitude"],
            "Vencedor": df.apply(_rotulo_vencedor, axis=1),
            "Percentual do vencedor": df["percentual_vencedor"],
            "Margem (p.p.)": df["margem_pp"],
            "Na frente no 2º turno": df["lider_disputa"] if "lider_disputa" in df.columns else "",
            "Margem da disputa (p.p.)": df["margem_disputa_pp"] if "margem_disputa_pp" in df.columns else None,
            "Votos em aberto": df["votos_outros"] if "votos_outros" in df.columns else None,
            "Votos válidos": df["votos_validos"],
            "Comparecimento": df["comparecimento"],
            "Votos computados no cargo": df["votos_computados"],
            "Seções": df["secoes"],
        }
    )


def _leia_me(ano: int, turno: int, cargo: str, limiar: float, meta: dict) -> pd.DataFrame:
    linhas = [
        ("Município", "Três Rios - RJ"),
        ("Ano", str(ano)),
        ("Turno", fmt_turno(turno)),
        ("Cargo", titulo_pt(cargo)),
        ("Gerado em", meta.get("gerado_em", "")),
        ("Origem", "Demonstração fictícia" if meta.get("origem") == "demonstracao" else "Arquivos oficiais do TSE"),
        ("Arquivos", meta.get("arquivos", "")),
        ("Votos válidos", "Nominais + legenda. Não incluem branco, nulo nem anulado."),
        ("Vencedor", "Candidato com mais votos nominais. Voto de legenda não define o vencedor."),
        ("Margem", "Diferença entre o 1º e o 2º, em pontos percentuais dos votos válidos."),
        ("Bairro prioritário", f"Margem menor que {limiar:g} pontos percentuais entre o 1º e o 2º do recorte."),
        (
            "2º turno",
            "Par de referência: os dois mais votados no país, quando o arquivo nacional foi processado. "
            "Em cada bairro, a margem da disputa usa só esses dois nomes.",
        ),
        (
            "Pode virar",
            "Os votos nominais dos demais candidatos são suficientes para cobrir a diferença entre os dois do par.",
        ),
        (
            "Prioritário no 2º turno",
            f"Margem da disputa abaixo de {limiar:g} pontos, ou bairro que ainda pode virar.",
        ),
        ("Abstenção", "Abstenções ÷ eleitores aptos das seções do bairro."),
        ("Mapa", "Cada ponto é um local de votação, não o contorno do bairro."),
    ]
    return pd.DataFrame(linhas, columns=["Campo", "Descrição"])


def _ajustar_planilha(aba) -> None:
    aba.freeze_panes = "A2"
    aba.auto_filter.ref = aba.dimensions
    for indice, coluna in enumerate(aba.columns, start=1):
        largura = 12
        for celula in list(coluna)[:40]:
            if celula.value is not None:
                largura = max(largura, min(len(str(celula.value)) + 2, 46))
        aba.column_dimensions[get_column_letter(indice)].width = largura
    for celula in aba[1]:
        celula.font = celula.font.copy(bold=True)


def gerar_excel(bairros: pd.DataFrame, candidatos: pd.DataFrame, locais: pd.DataFrame, meta: dict, ano: int, turno: int, cargo: str, limiar: float) -> bytes:
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as escritor:
        _leia_me(ano, turno, cargo, limiar, meta).to_excel(escritor, index=False, sheet_name="Leia-me")
        planilha_bairros = _bairros(bairros)
        planilha_bairros.to_excel(escritor, index=False, sheet_name="Bairros")
        if "Prioritário no 2º turno" in planilha_bairros.columns:
            prioritarios = planilha_bairros.loc[planilha_bairros["Prioritário no 2º turno"].eq("Sim")]
        else:
            prioritarios = planilha_bairros.loc[planilha_bairros["Prioritário"].eq("Sim")]
        prioritarios.to_excel(escritor, index=False, sheet_name="Prioritários")
        _candidatos(candidatos).to_excel(escritor, index=False, sheet_name="Candidatos")
        _locais(locais).to_excel(escritor, index=False, sheet_name="Locais")
        for nome in escritor.book.sheetnames:
            _ajustar_planilha(escritor.book[nome])
    return buffer.getvalue()


def gerar_csv(bairros: pd.DataFrame) -> bytes:
    """CSV com separador ';' e BOM, para abrir direto no Excel em português."""
    return _bairros(bairros).to_csv(sep=";", decimal=",", index=False).encode("utf-8-sig")
